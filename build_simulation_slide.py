"""Create a 13-slide deck and PDF by appending the Simulation Lab slide."""
from io import BytesIO
from pathlib import Path

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "RAYNEX One Sun, Five Jobs.pdf"
SHOT = ROOT / "raynex_ui_snapshot.png"
CROP = ROOT / "raynex_ui_snapshot_slide.png"
OUT_PPTX = ROOT / "RAYNEX One Sun, Five Jobs - Simulation.pptx"
OUT_PDF = ROOT / "RAYNEX One Sun, Five Jobs - Simulation.pdf"

W, H = 960, 540
GREEN = "193E2C"
MID_GREEN = "3F7D3A"
PALE = "F1F5E9"
CREAM = "FBF8EC"
GOLD = "EAAA22"
GOLD_PALE = "FBF0CF"
INK = "1C2A22"
MUTED = "5E6B63"
WHITE = "FFFFFF"
LINE = "D6DFD2"


def crop_ui() -> None:
    image = Image.open(SHOT).convert("RGB")
    target_h = int(image.width * 9 / 16)
    image.crop((0, 0, image.width, min(image.height, target_h))).save(CROP, quality=94)


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def add_rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x / 72), Inches(y / 72), Inches(w / 72), Inches(h / 72))
    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(fill)
    shape.line.color.rgb = rgb(line or fill)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def add_text(slide, text, x, y, w, h, size=12, color=INK, bold=False,
             font="Arial", align=PP_ALIGN.LEFT, margin=0, valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x / 72), Inches(y / 72), Inches(w / 72), Inches(h / 72))
    box.text_frame.clear(); box.text_frame.margin_left = box.text_frame.margin_right = Inches(margin / 72)
    box.text_frame.margin_top = box.text_frame.margin_bottom = 0
    box.text_frame.vertical_anchor = valign
    p = box.text_frame.paragraphs[0]; p.alignment = align; p.space_after = Pt(0)
    run = p.add_run(); run.text = text; run.font.name = font; run.font.size = Pt(size)
    run.font.bold = bold; run.font.color.rgb = rgb(color)
    return box


def add_rich_lines(slide, lines, x, y, w, h, size=10.5, color=INK):
    box = slide.shapes.add_textbox(Inches(x / 72), Inches(y / 72), Inches(w / 72), Inches(h / 72))
    tf = box.text_frame; tf.clear(); tf.margin_left = tf.margin_right = Inches(5 / 72)
    tf.margin_top = Inches(2 / 72); tf.word_wrap = True
    for i, (lead, rest) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(3.5); p.level = 0
        r = p.add_run(); r.text = lead; r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = True; r.font.color.rgb = rgb(GREEN)
        r = p.add_run(); r.text = rest; r.font.name = "Arial"; r.font.size = Pt(size); r.font.color.rgb = rgb(color)
    return box


def add_new_slide(prs: Presentation):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = rgb(PALE)

    add_text(slide, "Simulation Lab: Prove The Router Before The Hardware", 25, 13, 700, 28, 19, GREEN, True)
    add_text(slide, "A five-minute digital twin turns the architecture into a testable operating system.",
             25, 43, 720, 18, 10.5, INK)
    add_text(slide, "RAYNEX", 786, 10, 82, 18, 12, GREEN, True)
    add_text(slide, "13", 899, 4, 38, 28, 24, GREEN, True, align=PP_ALIGN.CENTER)
    add_text(slide, "of 13", 899, 34, 38, 14, 8, MUTED, True, align=PP_ALIGN.CENTER)
    add_rect(slide, 872, 50, 65, 4, "D5DECC")
    add_rect(slide, 872, 50, 65, 4, GOLD)

    add_rect(slide, 25, 67, 690, 31, GREEN)
    add_text(slide, "One screen validates energy priority, water logic, cold safety and farmer decisions.",
             35, 75, 670, 18, 11.5, WHITE, True)
    add_rect(slide, 720, 67, 217, 31, GOLD)
    add_text(slide, "D4: Prototype · E5: Technical Approach", 729, 76, 200, 16, 9, GREEN, True)

    # Screenshot, framed like the exhibit panels used across the source deck.
    add_rect(slide, 25, 107, 606, 310, WHITE, LINE)
    slide.shapes.add_picture(str(CROP), Inches(30 / 72), Inches(112 / 72),
                             width=Inches(596 / 72), height=Inches(298 / 72))
    add_rect(slide, 35, 386, 166, 20, GREEN, GREEN, True)
    add_text(slide, "LIVE FASTAPI + WEBSOCKET UI", 43, 390, 150, 12, 7.8, WHITE, True)

    # Dense right-side evidence stack.
    panels = [
        (107, 95, "WHAT IT MODELS", [
            ("15 kWp array · ", "weather-shaped PV output"),
            ("Five jobs · ", "cold → water → ice → dryer → grid"),
            ("Physical state · ", "8 zones, probes, pump, 5 t room"),
        ]),
        (207, 95, "HOW TO TEST", [
            ("Time · ", "1×, 60×, 600× or 3000×"),
            ("Faults · ", "cloud, heat, outage, probe, open door"),
            ("People · ", "approve, skip or stop the pump"),
        ]),
        (307, 110, "WHAT IT PROVES", [
            ("Priority · ", "critical cold load always comes first"),
            ("Safety · ", "daylight water + ice-backed cooling"),
            ("Evidence · ", "exact kWh ledger and run report"),
        ]),
    ]
    for y, height, title, lines in panels:
        add_rect(slide, 641, y, 296, height, WHITE, LINE)
        add_rect(slide, 641, y, 296, 22, GREEN)
        add_text(slide, title, 651, y + 5, 270, 13, 9.2, WHITE, True)
        add_rich_lines(slide, lines, 648, y + 28, 280, height - 31, 9.1)

    # Bottom process strip fills the page and explains the actual validation journey.
    add_rect(slide, 25, 429, 912, 83, CREAM, GOLD)
    add_rect(slide, 25, 429, 130, 83, GOLD)
    add_text(slide, "D4 SOFTWARE\nPROTOTYPE", 35, 443, 110, 34, 11, GREEN, True)
    add_text(slide, "Before hardware", 35, 483, 110, 14, 8.5, GREEN, True)
    steps = [
        ("1", "INJECT", "Change weather or break a sensor"),
        ("2", "OBSERVE", "Watch power, soil and cold state"),
        ("3", "DECIDE", "Approve or override the action"),
        ("4", "VERIFY", "Export evidence against targets"),
    ]
    sw = (782 / 4)
    for i, (num, title, desc) in enumerate(steps):
        x = 155 + i * sw
        if i: add_rect(slide, x, 437, 1, 67, LINE)
        add_text(slide, num, x + 10, 440, 25, 23, 15, GOLD, True)
        add_text(slide, title, x + 37, 443, 135, 16, 9.5, GREEN, True)
        add_text(slide, desc, x + 11, 468, 169, 31, 8.5, INK)
    return slide


def make_pptx() -> None:
    src = fitz.open(SOURCE)
    prs = Presentation(); prs.slide_width = Inches(13.333333); prs.slide_height = Inches(7.5)
    # Remove the default slide if a template ever supplies one.
    while prs.slides:
        rid = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rid); del prs.slides._sldIdLst[0]
    for page in src:
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(BytesIO(pix.tobytes("png")), 0, 0, width=prs.slide_width, height=prs.slide_height)
    add_new_slide(prs)
    prs.save(OUT_PPTX)


def pdf_text(page, text, rect, size=12, color=INK, bold=False, align=0):
    page.insert_textbox(fitz.Rect(*rect), text, fontsize=size, fontname="hebo" if bold else "helv",
                        color=tuple(int(color[i:i+2], 16) / 255 for i in (0, 2, 4)), align=align)


def pdf_line(page, text, x, baseline, size=12, color=INK, bold=False):
    page.insert_text((x, baseline), text, fontsize=size, fontname="hebo" if bold else "helv",
                     color=tuple(int(color[i:i+2], 16) / 255 for i in (0, 2, 4)))


def pdf_rect(page, rect, fill, line=None, width=.7):
    def c(v): return tuple(int(v[i:i+2], 16) / 255 for i in (0, 2, 4))
    page.draw_rect(fitz.Rect(*rect), color=c(line or fill), fill=c(fill), width=width)


def make_pdf() -> None:
    source = fitz.open(SOURCE); out = fitz.open(); out.insert_pdf(source)
    page = out.new_page(width=W, height=H)
    pdf_rect(page, (0, 0, W, H), PALE)
    pdf_line(page, "Simulation Lab: Prove The Router Before The Hardware", 25, 32, 19, GREEN, True)
    pdf_line(page, "A five-minute digital twin turns the architecture into a testable operating system.", 25, 56, 10.5)
    pdf_line(page, "RAYNEX", 786, 25, 12, GREEN, True)
    pdf_line(page, "13", 902, 28, 24, GREEN, True); pdf_line(page, "of 13", 909, 44, 8, MUTED, True)
    pdf_rect(page, (872, 50, 937, 54), GOLD)
    pdf_rect(page, (25, 67, 715, 98), GREEN); pdf_line(page, "One screen validates energy priority, water logic, cold safety and farmer decisions.", 35, 87, 11.5, WHITE, True)
    pdf_rect(page, (720, 67, 937, 98), GOLD); pdf_line(page, "D4: Prototype · E5: Technical Approach", 729, 87, 9, GREEN, True)
    pdf_rect(page, (25, 107, 631, 417), WHITE, LINE)
    page.insert_image(fitz.Rect(30, 112, 626, 410), filename=str(CROP))
    pdf_rect(page, (35, 386, 201, 406), GREEN); pdf_line(page, "LIVE FASTAPI + WEBSOCKET UI", 43, 400, 7.8, WHITE, True)
    panels = [
        (107, 202, "WHAT IT MODELS", ["15 kWp array - weather-shaped PV output", "Five jobs - cold > water > ice > dryer > grid", "Physical state - 8 zones, probes, pump, 5 t room"]),
        (207, 302, "HOW TO TEST", ["Time - 1x, 60x, 600x or 3000x", "Faults - cloud, heat, outage, probe, open door", "People - approve, skip or stop the pump"]),
        (307, 417, "WHAT IT PROVES", ["Priority - critical cold load always comes first", "Safety - daylight water + ice-backed cooling", "Evidence - exact kWh ledger and run report"]),
    ]
    for y1, y2, title, lines in panels:
        pdf_rect(page, (641, y1, 937, y2), WHITE, LINE); pdf_rect(page, (641, y1, 937, y1 + 22), GREEN)
        pdf_line(page, title, 651, y1 + 15, 9.2, WHITE, True)
        yy = y1 + 31
        for line in lines:
            pdf_text(page, "- " + line, (650, yy, 925, yy + 19), 8.8, INK); yy += 20
    pdf_rect(page, (25, 429, 937, 512), CREAM, GOLD); pdf_rect(page, (25, 429, 155, 512), GOLD)
    pdf_text(page, "D4 SOFTWARE\nPROTOTYPE", (35, 442, 145, 480), 11, GREEN, True); pdf_text(page, "Before hardware", (35, 484, 145, 500), 8.5, GREEN, True)
    steps = [("1", "INJECT", "Change weather or\nbreak a sensor"), ("2", "OBSERVE", "Watch power, soil\nand cold state"), ("3", "DECIDE", "Approve or override\nthe action"), ("4", "VERIFY", "Export evidence\nagainst targets")]
    sw = 782 / 4
    for i, (num, title, desc) in enumerate(steps):
        x = 155 + i * sw
        if i: pdf_rect(page, (x, 437, x + 1, 504), LINE)
        pdf_line(page, num, x + 10, 457, 15, GOLD, True)
        pdf_line(page, title, x + 37, 454, 9.5, GREEN, True)
        pdf_text(page, desc, (x + 11, 468, x + 180, 503), 8.5, INK)
    out.save(OUT_PDF, garbage=4, deflate=True)


if __name__ == "__main__":
    crop_ui(); make_pptx(); make_pdf()
    print(OUT_PPTX.name, OUT_PPTX.stat().st_size)
    print(OUT_PDF.name, OUT_PDF.stat().st_size)
